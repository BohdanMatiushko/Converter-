"""Excel export functionality for the Electric Units Converter."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = ["Значення", "Звідки", "Куди", "Результат", "Статус"]


def has_exportable_data(rows: Iterable[Sequence[str]]) -> bool:
    """Return True if there is at least one row with any meaningful data."""
    for row in rows:
        if any(str(cell).strip() for cell in row):
            return True
    return False


def export_to_excel(file_path: str | Path, rows: list[list[str]]) -> None:
    """
    Export rows to an Excel file.

    Args:
        file_path: Destination .xlsx path.
        rows: Table data rows.

    Raises:
        ValueError: If there is no data to export.
    """
    if not has_exportable_data(rows):
        raise ValueError("Немає даних для експорту")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Конвертація"

    header_fill = PatternFill(fill_type="solid", start_color="D9EAF7", end_color="D9EAF7")
    thin_side = Side(style="thin", color="C7C7C7")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col_index, header in enumerate(HEADERS, start=1):
        cell = worksheet.cell(row=1, column=col_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_index, row_data in enumerate(rows, start=2):
        for col_index, value in enumerate(row_data, start=1):
            cell = worksheet.cell(row=row_index, column=col_index, value=value)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:E{max(len(rows) + 1, 2)}"

    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            text = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(text))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 40)

    workbook.save(str(file_path))